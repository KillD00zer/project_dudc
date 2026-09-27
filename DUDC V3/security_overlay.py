"""
Cadastral Survey Certificate System - Top-Most Security Overlay & Anti-Forgery Watermark
========================================================================================
Protects certificate against cropping, scanner-forgery, and signature theft:
1. Generates high-resolution transparent A4 security watermark image.
2. Injects protective interlocking micro-lines directly over the bottom signature block.
3. Dynamically binds citizen name and DUDC encrypted security token into the watermark.
4. Anchors image in Word Header as top-most floating overlay (behindDoc="0") covering the whole page.
"""

import os
import docx
from docx.shared import Inches, Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from PIL import Image, ImageDraw, ImageFont

try:
    import arabic_reshaper
    from bidi.algorithm import get_display
    def shape_ar(text):
        reshaped = arabic_reshaper.reshape(str(text))
        return get_display(reshaped)
except Exception:
    def shape_ar(text):
        return str(text)

def get_font(size=24):
    for fn in ["arial.ttf", "tahoma.ttf", "calibri.ttf"]:
        try:
            return ImageFont.truetype(fn, size)
        except Exception:
            pass
    return ImageFont.load_default()

def generate_security_overlay_image(applicant_name: str, security_token: str, output_png_path: str) -> str:
    """
    Renders high-resolution transparent A4 overlay image (200 DPI = 1654x2338 px).
    Top-most overlay with light-gray alpha protecting text & signatures.
    """
    w, h = 1654, 2338  # Standard A4 at 200 DPI
    img = Image.new('RGBA', (w, h), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)

    f_body = get_font(28)
    f_token = get_font(22)
    f_micro = get_font(18)

    # 1. Repeating background watermark across certificate body
    color_body = (120, 120, 120, 75)
    color_token = (110, 110, 110, 80)

    txt_center = shape_ar("محافظة الدقهلية — مركز معلومات شبكات المرافق DUDC")
    txt_sub = shape_ar(f"وثيقة مساحية مؤمنة | {applicant_name}") + f" | {security_token}"

    for y in range(160, 1850, 240):
        draw.text((120, y), txt_center, font=f_body, fill=color_body)
        draw.text((200, y + 80), txt_sub, font=f_token, fill=color_token)

    # 2. Signature Guard Area (y: 1900 to 2320 px)
    # Directly overlays the 'يعتمد' line and the 4 signature columns:
    # (فنى مساحة / مسئول النظم / رئيس قسم النظم / مدير المركز)
    color_sig_lines = (110, 110, 110, 90)
    color_sig_text = (100, 100, 100, 95)
    
    # Interlocking security micro-lines cutting across signatures
    for y_line in range(1920, 2320, 36):
        draw.line([(60, y_line), (w - 60, y_line)], fill=color_sig_lines, width=1)
        sig_label = shape_ar(f"-- ختم التوقيع المؤمن -- {applicant_name}") + f" -- {security_token} --"
        draw.text((80, y_line - 15), sig_label, font=f_micro, fill=color_sig_text)

    os.makedirs(os.path.dirname(os.path.abspath(output_png_path)), exist_ok=True)
    img.save(output_png_path, 'PNG')
    return output_png_path

def apply_watermark_overlay_to_docx(doc: docx.Document, overlay_png_path: str):
    """
    Injects overlay PNG into Word Document header with behindDoc='0'
    so it floats OVER all document elements and signatures as a protective top layer.
    """
    section = doc.sections[0]
    header = section.header
    
    p = header.paragraphs[0]
    p.text = ""  # clear any default text
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    
    run = p.add_run()
    run.add_picture(overlay_png_path, width=Inches(8.27), height=Inches(11.69))

    inline_elem = run._r.xpath('.//wp:inline')[0]
    extent = inline_elem.find('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent')
    docPr = inline_elem.find('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr')
    cNvGraphicFramePr = inline_elem.find('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}cNvGraphicFramePr')
    graphic = inline_elem.find('{http://schemas.openxmlformats.org/drawingml/2006/main}graphic')

    cx = extent.get('cx')
    cy = extent.get('cy')

    # behindDoc="0" places it in front of document text and signatures!
    anchor_xml = f'''<wp:anchor {nsdecls('wp')} distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251658240" behindDoc="0" locked="0" layoutInCell="0" allowOverlap="1">
        <wp:simplePos x="0" y="0"/>
        <wp:positionH relativeFrom="page">
            <wp:posOffset>0</wp:posOffset>
        </wp:positionH>
        <wp:positionV relativeFrom="page">
            <wp:posOffset>0</wp:posOffset>
        </wp:positionV>
        <wp:extent cx="{cx}" cy="{cy}"/>
        <wp:wrapNone/>
    </wp:anchor>'''
    
    anchor_elem = parse_xml(anchor_xml)
    anchor_elem.append(docPr)
    anchor_elem.append(cNvGraphicFramePr)
    anchor_elem.append(graphic)

    drawing = inline_elem.getparent()
    drawing.remove(inline_elem)
    drawing.append(anchor_elem)
