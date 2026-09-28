"""
Cadastral Survey Certificate System - Word Document Automation & Publisher
==========================================================================
Populates the official 7-column template 'شهادة.docx':
1. Fills applicant metadata into Table 0 (Rows 1-5) using Arial 12.
2. Formats all headlines and headers using Arial 14 Bold.
3. Populates 4 Cardinal Boundaries (البحري، الشرقي، القبلي، الغربي) with Vertical Merge:
   - Each boundary takes a large merged cell for boundary name, neighbor description, and total length.
   - Across from it are individual sub-rows with the exact coordinates (LONG & LAT) for each vertex belonging to that boundary.
   - Clean neighbor descriptions without dimension summation formulas like (89.04+31.76).
   - Formats coordinates in WGS84 Decimal Degrees (DD, exactly 8 decimals: .8f) using Arial 12 Bold.
4. Embeds CAD Croquis and Google Satellite images into the final row.
5. Updates the official employee signature block with selected technician & GIS officer using Arial 12 Bold.
6. Enforces strict single-page A4 layout constraints with zero paragraph margins.
"""

import os
import copy
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from datetime import datetime
import security_overlay

ARABIC_DAYS = ["الإثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"]

def to_arabic_numerals(text):
    if text is None:
        return ""
    trans = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")
    return str(text).translate(trans)

def format_issue_date(issue_date_val=None):
    if isinstance(issue_date_val, str) and issue_date_val.strip():
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d-%m-%Y", "%d/%m/%Y"):
            try:
                dt = datetime.strptime(issue_date_val.strip(), fmt)
                break
            except ValueError:
                dt = datetime.now()
    elif isinstance(issue_date_val, datetime):
        dt = issue_date_val
    else:
        dt = datetime.now()
    day_name = ARABIC_DAYS[dt.weekday()]
    date_str = to_arabic_numerals(f"{dt.year:04d}/{dt.month:02d}/{dt.day:02d}")
    display_text = f"تحريراً في : {day_name} الموافق {date_str}"
    iso_date = f"{dt.year:04d}-{dt.month:02d}-{dt.day:02d}"
    return display_text, iso_date

import sys
if getattr(sys, 'frozen', False):
    _APP_DIR = os.path.dirname(os.path.abspath(sys.executable))
    _BUNDLE_DIR = getattr(sys, '_MEIPASS', _APP_DIR)
else:
    _APP_DIR = os.path.dirname(os.path.abspath(__file__))
    _BUNDLE_DIR = _APP_DIR

_p_app = os.path.join(_APP_DIR, "شهادة.docx")
_p_bundle = os.path.join(_BUNDLE_DIR, "شهادة.docx")
DEFAULT_TEMPLATE = _p_app if os.path.exists(_p_app) else _p_bundle

def format_coord_dd(val):
    """
    Strictly formats coordinates to 8 decimal places (WGS84 DD precision).
    Ensures all 8 numbers after the decimal point are preserved without truncation.
    """
    try:
        f = float(val)
        return f"{f:.8f}"
    except (ValueError, TypeError):
        return str(val)

def apply_run_font(run, font_name="Arial", size_pt=12, bold=None):
    """
    Applies font family, point size, and bold formatting to a run.
    Explicitly sets Word XML rFonts (ascii, hAnsi, cs) to ensure Arial renders
    correctly in Arabic and complex script environments.
    """
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
        
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(docx.oxml.ns.qn('w:ascii'), font_name)
    rFonts.set(docx.oxml.ns.qn('w:hAnsi'), font_name)
    rFonts.set(docx.oxml.ns.qn('w:cs'), font_name)

    if bold:
        bCs = rPr.find(docx.oxml.ns.qn('w:bCs'))
        if bCs is None:
            rPr.append(docx.oxml.OxmlElement('w:bCs'))
    rtl = rPr.find(docx.oxml.ns.qn('w:rtl'))
    if rtl is None:
        rPr.append(docx.oxml.OxmlElement('w:rtl'))

def set_cell_text(cell, text, bold=False, font_size=12, font_color=None, align=WD_ALIGN_PARAGRAPH.CENTER):
    """
    Clears cell content and sets new text with strict Arial typography,
    zero margins, and tight 1.0 line spacing.
    Default font size is Arial 12.
    """
    p = cell.paragraphs[0]
    p.text = str(text)
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    
    if p.runs:
        run = p.runs[0]
        apply_run_font(run, font_name="Arial", size_pt=font_size, bold=bold)
        if font_color:
            run.font.color.rgb = font_color
        return run
    return None

def update_signature_names(doc, survey_tech="محمد ابراهيم بدير", sys_officer="شريف محمد"):
    """
    Updates the employee signature blocks at the bottom of the certificate
    with the surveyor's and GIS officer's official names using Arial 12 Bold.
    Maintains exact 4-column spacing across the line to ensure strict 1-page fit.
    """
    col1 = f"أ / {str(survey_tech).strip()}"
    gap1 = " " * max(4, 35 - len(col1))
    col2 = f"أ / {str(sys_officer).strip()}"
    cur_len = len(col1) + len(gap1) + len(col2)
    gap2 = " " * max(4, 65 - cur_len)
    col3 = "م / محمد مصطفى"
    cur_len2 = cur_len + len(gap2) + len(col3)
    gap3 = " " * max(4, 95 - cur_len2)
    col4 = "م / السيد زين العابدين"
    new_sig_line = f"{col1}{gap1}{col2}{gap2}{col3}{gap3}{col4}"

    # 1. Primary: Search in document paragraphs (where official template places signatures)
    updated = False
    for p in doc.paragraphs:
        txt = p.text
        if any(k in txt for k in ["السيد زين العابدين", "محمد مصطفى", "ابراهيم بدير"]):
            p.text = new_sig_line
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if p.runs:
                apply_run_font(p.runs[0], "Arial", 12, bold=True)
            updated = True
            break

    # 2. Fallback: Search in tables in case template is ever converted to table format
    if not updated:
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        txt = p.text
                        if any(k in txt for k in ["السيد زين العابدين", "محمد مصطفى", "ابراهيم بدير"]):
                            p.text = new_sig_line
                            p.paragraph_format.space_before = Pt(0)
                            p.paragraph_format.space_after = Pt(0)
                            p.paragraph_format.line_spacing = 1.0
                            if p.runs:
                                apply_run_font(p.runs[0], "Arial", 12, bold=True)
                            updated = True
                            break
                        elif "فني المساحة" in txt or "فني مساحة" in txt:
                            p.text = f"فني مساحة / {survey_tech}"
                            p.paragraph_format.space_before = Pt(0)
                            p.paragraph_format.space_after = Pt(0)
                            p.paragraph_format.line_spacing = 1.0
                            if p.runs:
                                apply_run_font(p.runs[0], "Arial", 12, bold=True)
                        elif "مسؤول النظم" in txt or "مسئول النظم" in txt:
                            p.text = f"مسئول النظم / {sys_officer}"
                            p.paragraph_format.space_before = Pt(0)
                            p.paragraph_format.space_after = Pt(0)
                            p.paragraph_format.line_spacing = 1.0
                            if p.runs:
                                apply_run_font(p.runs[0], "Arial", 12, bold=True)

def format_all_document_typography(doc):
    """
    Enforces Arial 14 Bold for headlines and Arial 12 for metadata and boundary data.
    """
    if not doc.tables:
        return
        
    table = doc.tables[0]
    
    # Row 0: 'بيانات مقدم الطلب' -> Arial 14 Bold
    if len(table.rows) > 0:
        for cell in table.rows[0].cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    apply_run_font(r, "Arial", 14, bold=True)
                    
    # Row 6: 'الحدود والابعاد' | 'إحداثيات الموقع' -> Arial 14 Bold
    if len(table.rows) > 6:
        for cell in table.rows[6].cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    apply_run_font(r, "Arial", 14, bold=True)
                    
    # Row 7: 'LONG (E)' | 'LAT (N)' -> Arial 12 Bold
    if len(table.rows) > 7:
        for cell in table.rows[7].cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    apply_run_font(r, "Arial", 12, bold=True)

def generate_certificate_docx(parcel, croquis_img_path, satellite_img_path, output_docx_path, template_path=None, survey_tech="محمد ابراهيم بدير", sys_officer="شريف محمد", issue_date_val=None, security_token=None):
    """
    Generates the official Cadastral Survey Certificate Word document (.docx).
    Each boundary has a large vertically merged cell on the left, and across from it
    are sub-rows containing the exact coordinates (LONG, LAT) of each vertex belonging to it.
    Clean neighbor descriptions without dimension summation formulas.
    """
    t_path = template_path or DEFAULT_TEMPLATE
    if not os.path.exists(t_path):
        raise FileNotFoundError(f"Certificate template not found at: {t_path}")
        
    doc = docx.Document(t_path)
    table = doc.tables[0]
    
    # Ensure tight document page margins
    section = doc.sections[0]
    section.top_margin = Inches(0.20)
    section.bottom_margin = Inches(0.10)
    section.left_margin = Inches(0.50)
    section.right_margin = Inches(0.50)
    
    # Ensure table alignment is centered
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Inject official issue date right under 'مركز معلومات شبكات المرافق'
    date_display, _ = format_issue_date(issue_date_val)
    if len(doc.paragraphs) > 2:
        p_date = doc.paragraphs[2]
        p_date.text = date_display
        p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_date.paragraph_format.space_before = Pt(2)
        p_date.paragraph_format.space_after = Pt(2)
        p_date.paragraph_format.line_spacing = 1.0
        if p_date.runs:
            run = p_date.runs[0]
            apply_run_font(run, font_name="Arial", size_pt=12, bold=True)
            run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    
    # 1. Applicant & Parcel Metadata (Rows 1 to 5)
    # Row 1: Applicant Name & Receipt No
    set_cell_text(table.rows[1].cells[1], parcel.get("applicant_name", ""), bold=True, font_size=12)
    set_cell_text(table.rows[1].cells[5], to_arabic_numerals(parcel.get("receipt_no", "")), bold=True, font_size=12)
    
    # Row 2: National ID & Request No
    set_cell_text(table.rows[2].cells[1], to_arabic_numerals(parcel.get("national_id", "")), bold=True, font_size=12)
    set_cell_text(table.rows[2].cells[5], to_arabic_numerals(parcel.get("request_no", "")), bold=True, font_size=12)
    
    # Row 3: District & Village
    set_cell_text(table.rows[3].cells[1], parcel.get("district", ""), bold=True, font_size=12)
    village_str = parcel.get("village", "") or "--------"
    set_cell_text(table.rows[3].cells[5], village_str, bold=True, font_size=12)
    
    # Row 4: Address & Area
    set_cell_text(table.rows[4].cells[1], parcel.get("address", ""), bold=True, font_size=12)
    area_val = parcel.get('stated_area_m2', '')
    area_str = f"{to_arabic_numerals(area_val)} م²" if area_val else ""
    set_cell_text(table.rows[4].cells[5], area_str, bold=True, font_size=12)
    
    # Row 5: Transaction & Site Status
    set_cell_text(table.rows[5].cells[1], parcel.get("transaction_type", "إنشاء"), bold=False, font_size=12)
    set_cell_text(table.rows[5].cells[5], parcel.get("site_status", "أرض فضاء"), bold=False, font_size=12)
    
    # 2. 4 Cardinal Boundaries with Vertical Merging
    cardinal_dirs = [
        {"key": "north", "name": "الحد البحري"},
        {"key": "east",  "name": "الحد الشرقي"},
        {"key": "south", "name": "الحد القبلي"},
        {"key": "west",  "name": "الحد الغربي"},
    ]
    
    verts = parcel["vertices"]
    segments = parcel.get("segments", [])
    n_pts = len(verts)
    bounds = parcel.get("boundaries", {})
    
    # Pre-map any missing segment directions
    lons = [v["lon"] for v in verts]
    lats = [v["lat"] for v in verts]
    c_lon = float(sum(lons)) / len(lons) if lons else 0.0
    c_lat = float(sum(lats)) / len(lats) if lats else 0.0
    
    for i, seg in enumerate(segments):
        if not seg.get("direction"):
            p1_v = next((v for v in verts if v["point_index"] == seg["from_point"]), verts[i % n_pts])
            p2_v = next((v for v in verts if v["point_index"] == seg["to_point"]), verts[(i + 1) % n_pts])
            dx = p2_v["lon"] - p1_v["lon"]
            dy = p2_v["lat"] - p1_v["lat"]
            mid_x = (p1_v["lon"] + p2_v["lon"]) / 2.0 - c_lon
            mid_y = (p1_v["lat"] + p2_v["lat"]) / 2.0 - c_lat
            seg["direction"] = ("north" if mid_y > 0 else "south") if abs(dx) > abs(dy) else ("east" if mid_x > 0 else "west")

    # Group segments and vertices by boundary
    dir_data = []
    for dir_idx, d in enumerate(cardinal_dirs):
        matching_segs = [s for s in segments if s.get("direction") == d["key"]]
        if not matching_segs:
            # Fallback if direction has no segments
            seg_pts = [verts[dir_idx % n_pts]]
            total_len = 0.0
        else:
            seg_pts = []
            for s in matching_segs:
                pt = next((v for v in verts if v["point_index"] == s["from_point"]), None)
                if pt and pt not in seg_pts:
                    seg_pts.append(pt)
            if not seg_pts:
                seg_pts = [next((v for v in verts if v["point_index"] == matching_segs[0]["from_point"]), verts[0])]
            total_len = sum(s.get("length_m", 0.0) for s in matching_segs)
            
        # Clean neighbor description WITHOUT summation formula string
        clean_neighbor = bounds.get(d["key"], "").strip()
        dir_data.append({
            "dir": d,
            "total_len": total_len,
            "neighbor": clean_neighbor,
            "points": seg_pts
        })

    # Prepare table XML replacement for boundary rows (starting at row 8)
    ref_tr = table.rows[8]._tr
    tbl_elem = table._tbl
    row12_elem = table.rows[12]._tr

    # Remove template rows 8..11
    for r_i in range(11, 7, -1):
        tbl_elem.remove(table.rows[r_i]._tr)

    # Insert grouped rows with vMerge before row12
    for item in dir_data:
        pts = item["points"]
        num_pts = len(pts)
        for p_idx in range(num_pts):
            new_tr = copy.deepcopy(ref_tr)
            tcs = new_tr.xpath('./w:tc')
            
            # Apply vertical merge on first 4 columns if multiple points
            if num_pts > 1:
                vmerge_val = 'restart' if p_idx == 0 else ''
                for c_i in range(4):
                    tcPr = tcs[c_i].get_or_add_tcPr()
                    for vm in tcPr.findall(qn('w:vMerge')):
                        tcPr.remove(vm)
                    if vmerge_val == 'restart':
                        tcPr.append(parse_xml(f'<w:vMerge {nsdecls("w")} w:val="restart"/>'))
                    else:
                        tcPr.append(parse_xml(f'<w:vMerge {nsdecls("w")}/>'))
                        for p_el in tcs[c_i].xpath('./w:p'):
                            p_el.getparent().remove(p_el)
                        tcs[c_i].append(parse_xml(f'<w:p {nsdecls("w")}/>'))
            
            row12_elem.addprevious(new_tr)

    # Save and reload to ensure clean python-docx row wrapper bindings
    doc.save(output_docx_path)
    doc = docx.Document(output_docx_path)
    table = doc.tables[0]

    # Populate cell contents
    curr_r = 8
    for item in dir_data:
        pts = item["points"]
        num_pts = len(pts)
        dir_name = item["dir"]["name"]
        neighbor = item["neighbor"]
        total_len_str = to_arabic_numerals(f"{item['total_len']:.2f}")
        
        # Set text on first row of boundary
        row0 = table.rows[curr_r]
        set_cell_text(row0.cells[0], dir_name, bold=True, font_size=12)
        set_cell_text(row0.cells[1], neighbor, bold=True if neighbor else False, font_size=12)
        set_cell_text(row0.cells[2], "بطول", bold=False, font_size=12)
        set_cell_text(row0.cells[3], total_len_str, bold=True, font_size=12)
        
        # Set coordinates for each point row
        for p_idx, pt in enumerate(pts):
            r = table.rows[curr_r + p_idx]
            lon_str = format_coord_dd(pt["lon"])
            lat_str = format_coord_dd(pt["lat"])
            set_cell_text(r.cells[4], lon_str, bold=True, font_size=12)
            set_cell_text(r.cells[6], lat_str, bold=True, font_size=12)
            
        curr_r += num_pts

    # Format Croquis and Satellite Header Row (the row right after boundary rows)
    header_row_idx = curr_r
    if len(table.rows) > header_row_idx:
        for c in table.rows[header_row_idx].cells:
            for p_el in c.paragraphs:
                for r_run in p_el.runs:
                    apply_run_font(r_run, "Arial", 14, bold=True)

    # Format all document typography (Arial 14 Bold headers, Arial 12 body)
    format_all_document_typography(doc)

    # 3. Image Insertion (Last Row)
    img_row_idx = len(table.rows) - 1
    img_row = table.rows[img_row_idx]
    
    # Left Cell (Cells 0-2): CAD Dimensions Croquis
    cell_croq = img_row.cells[0]
    p_croq = cell_croq.paragraphs[0]
    p_croq.text = ""
    p_croq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_croq.paragraph_format.space_before = Pt(0)
    p_croq.paragraph_format.space_after = Pt(0)
    p_croq.paragraph_format.line_spacing = 1.0
    
    if os.path.exists(croquis_img_path):
        run_croq = p_croq.add_run()
        run_croq.add_picture(croquis_img_path, width=Inches(3.05))

    # Right Cell (Cells 3-6): Google Satellite Imagery
    cell_sat = img_row.cells[3]
    p_sat = cell_sat.paragraphs[0]
    p_sat.text = ""
    p_sat.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sat.paragraph_format.space_before = Pt(0)
    p_sat.paragraph_format.space_after = Pt(0)
    p_sat.paragraph_format.line_spacing = 1.0
    
    if os.path.exists(satellite_img_path):
        run_sat = p_sat.add_run()
        run_sat.add_picture(satellite_img_path, width=Inches(3.05))

    # 4. Official Signatures (Survey Technician & GIS Officer)
    update_signature_names(doc, survey_tech=survey_tech, sys_officer=sys_officer)

    # 5. Security Token Verification Line in Certificate (Paragraph 9)
    if security_token and len(doc.paragraphs) > 9:
        p_token = doc.paragraphs[9]
        p_token.text = f"كود التأمين والتحقق الرقمي (DUDC Token): {security_token}"
        p_token.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_token.paragraph_format.space_before = Pt(0)
        p_token.paragraph_format.space_after = Pt(0)
        p_token.paragraph_format.line_spacing = 1.0
        if p_token.runs:
            run = p_token.runs[0]
            apply_run_font(run, font_name="Arial", size_pt=10, bold=True)
            run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # 6. Top-Most Protective Security Overlay (Watermark Guard over signatures & canvas)
    if security_token:
        applicant_name = parcel.get("applicant_name", "")
        pid = parcel.get("parcel_id", "cert")
        overlay_png_path = os.path.join(os.path.dirname(output_docx_path), f"overlay_{pid}.png")
        try:
            security_overlay.generate_security_overlay_image(applicant_name, security_token, overlay_png_path)
            security_overlay.apply_watermark_overlay_to_docx(doc, overlay_png_path)
            if os.path.exists(overlay_png_path):
                os.remove(overlay_png_path)
        except Exception as e:
            print(f"[!] Warning: Security overlay failed: {e}")

    # Save final publication document
    doc.save(output_docx_path)
    return output_docx_path
